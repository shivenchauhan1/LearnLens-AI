import io
import re
import math
import random
from collections import Counter
from typing import Dict, List, Tuple

import streamlit as st
from pypdf import PdfReader
from pptx import Presentation

st.set_page_config(page_title="LearnLens AI", page_icon="📚", layout="wide")

STOPWORDS = set("a an and are as at be been being by can could did do does for from had has have he her hers him his how i if in into is it its may me more most my no not of on or our ours she should so than that the their theirs them then there these they this those to too up us was we were what when where which who why will with would you your yours about after all also any because before between both but each few further here once only other own same some such through under until very while within without system topic course material lecture chapter page slide student learning concept process using used use example examples first second third according another different important information explain understanding".split())


def clean_text(s: str) -> str:
    return re.sub(r"\s+", " ", s or "").strip()


def extract_pdf(data: bytes, filename: str) -> List[dict]:
    reader = PdfReader(io.BytesIO(data))
    units = []
    for i, page in enumerate(reader.pages, start=1):
        text = clean_text(page.extract_text() or "")
        if text:
            units.append({"source": filename, "location": f"Page {i}", "kind": "PDF", "text": text})
    return units


def extract_pptx(data: bytes, filename: str) -> List[dict]:
    prs = Presentation(io.BytesIO(data))
    units = []
    for i, slide in enumerate(prs.slides, start=1):
        bits = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = clean_text(shape.text)
                if text:
                    bits.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    bits.append(" | ".join(clean_text(cell.text) for cell in row.cells))
        text = clean_text(". ".join(bits))
        if text:
            units.append({"source": filename, "location": f"Slide {i}", "kind": "PPTX", "text": text})
    return units


def parse_transcript(data: bytes, filename: str) -> List[dict]:
    raw = data.decode("utf-8", errors="ignore").replace("\r", "")
    blocks = re.split(r"\n\s*\n", raw)
    units = []
    for block in blocks:
        lines = [x.strip() for x in block.splitlines() if x.strip()]
        if not lines:
            continue
        timestamp = next((x for x in lines if "-->" in x), None)
        content_lines = [x for x in lines if x != timestamp and not x.isdigit()]
        text = clean_text(" ".join(content_lines))
        if text:
            loc = f"Timestamp {timestamp.split('-->')[0].strip()}" if timestamp else "Transcript segment"
            units.append({"source": filename, "location": loc, "kind": "Transcript", "text": text})
    if not units:
        text = clean_text(raw)
        if text:
            units = [{"source": filename, "location": "Transcript", "kind": "Text", "text": text}]
    return units


def tokenize(text: str) -> List[str]:
    return re.findall(r"[a-zA-Z0-9][a-zA-Z0-9_+-]*", text.lower())


def split_units(units: List[dict], max_words: int = 180) -> List[dict]:
    chunks = []
    for unit in units:
        # Keep the original page/slide/timestamp as the citation anchor.
        sentences = re.split(r"(?<=[.!?])\s+", unit["text"])
        current = []
        size = 0
        for sentence in sentences:
            words = sentence.split()
            if current and size + len(words) > max_words:
                chunk = dict(unit)
                chunk["text"] = " ".join(current)
                chunks.append(chunk)
                current, size = [], 0
            if words:
                current.append(sentence)
                size += len(words)
        if current:
            chunk = dict(unit)
            chunk["text"] = " ".join(current)
            chunks.append(chunk)
    return chunks


def retrieve(query: str, chunks: List[dict], limit: int = 3) -> List[Tuple[float, dict]]:
    q_tokens = [t for t in tokenize(query) if t not in STOPWORDS]
    if not q_tokens:
        q_tokens = tokenize(query)
    docs = [set(tokenize(c["text"])) for c in chunks]
    df = Counter(t for doc in docs for t in doc)
    n = max(1, len(docs))
    scored = []
    for chunk, terms in zip(chunks, docs):
        counts = Counter(tokenize(chunk["text"]))
        score = 0.0
        for term in q_tokens:
            if term in terms:
                score += (1 + math.log(1 + counts[term])) * math.log(1 + (n / (1 + df[term])))
        if score:
            scored.append((score, chunk))
    scored.sort(key=lambda x: x[0], reverse=True)
    return scored[:limit]


def make_quiz(chunks: List[dict], used: set, count: int = 5) -> List[dict]:
    questions = []
    seen_now = set(used)
    for chunk in chunks:
        # Heuristic cloze questions are intentionally labelled as baseline-generated.
        sentences = re.split(r"(?<=[.!?])\s+", chunk["text"])
        for sentence in sentences:
            sentence = clean_text(sentence)
            words = re.findall(r"\b[A-Za-z][A-Za-z0-9_-]{4,}\b", sentence)
            candidates = [w for w in words if w.lower() not in STOPWORDS]
            if len(sentence.split()) < 7 or not candidates:
                continue
            # Prefer a distinctive longer term, but skip sentences with weak candidates.
            candidates.sort(key=lambda w: (len(w), w.lower()), reverse=True)
            answer = candidates[0].strip(".,:;!?()[]{}\"'")
            key = re.sub(r"[^a-z0-9]", "", sentence.lower())
            if key in seen_now or len(answer) < 5:
                continue
            question_text = re.sub(r"\b" + re.escape(answer) + r"\b", "________", sentence, count=1, flags=re.IGNORECASE)
            if question_text == sentence:
                continue
            seen_now.add(key)
            questions.append({"question": f"Complete the statement: {question_text}", "answer": answer, "source": chunk["source"], "location": chunk["location"], "topic": chunk["source"], "evidence": sentence})
    random.shuffle(questions)
    return questions[:count]


st.markdown("<div style='padding:8px 0 2px 0'><span style='font-size:2.25rem;font-weight:750'>📚 LearnLens AI</span><br><span style='color:#8795a7'>A source-aware study companion · Track D prototype</span></div>", unsafe_allow_html=True)
st.caption("Prototype mode: answers are retrieved from uploaded content; quiz questions use transparent extractive cloze generation. No generative AI claims are made.")

with st.sidebar:
    st.header("1 · Add course material")
    files = st.file_uploader("Upload PDFs, PowerPoint slides, or transcript files", type=["pdf", "pptx", "txt", "srt", "vtt"], accept_multiple_files=True)
    st.info("For lecture videos, upload an existing .srt/.vtt transcript. Direct video transcription and diagram OCR are not implemented in this MVP.")
    st.divider()
    st.header("Study settings")
    quiz_length = st.slider("Questions per quiz", min_value=3, max_value=10, value=5)

if "mastery" not in st.session_state:
    st.session_state.mastery = {}
if "used_questions" not in st.session_state:
    st.session_state.used_questions = set()
if "quiz" not in st.session_state:
    st.session_state.quiz = []
if "quiz_feedback" not in st.session_state:
    st.session_state.quiz_feedback = None
if "quiz_index" not in st.session_state:
    st.session_state.quiz_index = 0

raw_units = []
if files:
    for uploaded in files:
        try:
            data = uploaded.getvalue()
            suffix = uploaded.name.lower().rsplit(".", 1)[-1]
            if suffix == "pdf":
                raw_units.extend(extract_pdf(data, uploaded.name))
            elif suffix == "pptx":
                raw_units.extend(extract_pptx(data, uploaded.name))
            else:
                raw_units.extend(parse_transcript(data, uploaded.name))
        except Exception as exc:
            st.error(f"Could not read {uploaded.name}: {exc}")

chunks = split_units(raw_units)

if not files:
    st.markdown("### Start with a course file")
    c1, c2, c3 = st.columns(3)
    c1.metric("Supported sources", "PDF · PPTX")
    c2.metric("Transcript sources", "SRT · VTT · TXT")
    c3.metric("Core workflow", "Cite · Quiz · Learn")
    st.write("Upload class notes, a textbook chapter, lecture slides, or a transcript in the sidebar. LearnLens extracts content and keeps each passage tied to its source location.")
    st.stop()

if not chunks:
    st.warning("No extractable text was found. The PDF may be scanned or the slides may contain image-only text. Try a text-based PDF/PPTX or a transcript file.")
    st.stop()

num_sources = len(set(c["source"] for c in chunks))
col1, col2, col3, col4 = st.columns(4)
col1.metric("Sources loaded", num_sources)
col2.metric("Indexed passages", len(chunks))
col3.metric("Quiz attempts", sum(x.get("attempts", 0) for x in st.session_state.mastery.values()))
col4.metric("Tracked topics", len(st.session_state.mastery))

ask_tab, quiz_tab, progress_tab, sources_tab = st.tabs(["💬 Ask your material", "🧠 Adaptive quiz", "📈 Learning progress", "📑 Source library"])

with ask_tab:
    st.subheader("Ask a question")
    query = st.text_input("Question", placeholder="e.g., What is a system call?", key="tutor_query")
    top_n = st.slider("Evidence passages", 1, 5, 3, key="evidence_n")
    if st.button("Find source-backed answer", type="primary", key="ask_button"):
        hits = retrieve(query, chunks, top_n) if query.strip() else []
        st.session_state.last_hits = hits
        st.session_state.last_query = query
    if "last_hits" in st.session_state and st.session_state.get("last_query"):
        hits = st.session_state.last_hits
        if not hits:
            st.warning("I couldn't find supporting text in the files you uploaded. Try different wording or add course material that covers this question. I won't invent a source-backed answer.")
        else:
            st.success(f"Found {len(hits)} relevant passage(s) in your uploaded material.")
            st.markdown("**Evidence-based response**")
            st.write("The passages below are the best matches to your question. Use the citations to verify the exact course context; this MVP retrieves evidence rather than generating a free-form LLM explanation.")
            for i, (score, hit) in enumerate(hits, start=1):
                with st.expander(f"Source {i}: {hit['source']} · {hit['location']}", expanded=(i == 1)):
                    st.write(hit["text"])
                    st.caption(f"Citation: {hit['source']} — {hit['location']} · Retrieval score: {score:.2f}")

with quiz_tab:
    st.subheader("Check what you know")
    st.write("Questions are generated from extracted course sentences. Correctness is checked against the source term; this is a lightweight baseline, not a validated exam generator.")
    if st.button("Generate a new quiz", type="primary"):
        st.session_state.quiz = make_quiz(chunks, st.session_state.used_questions, quiz_length)
        st.session_state.quiz_index = 0
        st.session_state.quiz_feedback = None
    quiz = st.session_state.quiz
    if quiz:
        idx = st.session_state.quiz_index
        if idx >= len(quiz):
            st.success("Quiz complete. Review the Learning progress tab for your topic-level results.")
            if st.button("Start another quiz"):
                st.session_state.quiz = make_quiz(chunks, st.session_state.used_questions, quiz_length)
                st.session_state.quiz_index = 0
                st.session_state.quiz_feedback = None
                st.rerun()
        else:
            q = quiz[idx]
            st.progress(idx / max(1, len(quiz)), text=f"Question {idx + 1} of {len(quiz)}")
            st.markdown(f"#### {q['question']}")
            answer_input = st.text_input("Your answer", key=f"answer_{idx}")
            if st.button("Check answer", key=f"check_{idx}"):
                norm = lambda x: re.sub(r"[^a-z0-9]", "", x.lower().strip())
                correct = norm(answer_input) == norm(q["answer"])
                st.session_state.quiz_feedback = {"correct": correct, "answer": q["answer"], "evidence": q["evidence"], "source": q["source"], "location": q["location"], "topic": q["topic"]}
                st.session_state.used_questions.add(re.sub(r"[^a-z0-9]", "", q["evidence"].lower()))
                m = st.session_state.mastery.setdefault(q["topic"], {"correct": 0, "attempts": 0})
                m["attempts"] += 1
                if correct:
                    m["correct"] += 1
            feedback = st.session_state.quiz_feedback
            if feedback:
                if feedback["correct"]:
                    st.success("Correct — well done!")
                else:
                    st.warning(f"Not quite. Expected answer: **{feedback['answer']}**")
                st.caption(f"Source: {feedback['source']} — {feedback['location']}")
                st.write(f"Source statement: {feedback['evidence']}")
                if st.button("Next question", key=f"next_{idx}"):
                    st.session_state.quiz_index += 1
                    st.session_state.quiz_feedback = None
                    st.rerun()
    else:
        st.info("Generate a quiz to begin. Upload material with complete explanatory sentences for better questions.")

with progress_tab:
    st.subheader("Your learning progress")
    if st.session_state.mastery:
        rows = []
        for topic, values in st.session_state.mastery.items():
            attempts = values["attempts"]
            score = values["correct"] / attempts if attempts else 0
            rows.append({"Source/topic": topic, "Correct": values["correct"], "Attempts": attempts, "Observed accuracy": f"{score:.0%}", "Recommended action": "Review source material" if score < 0.7 else "Try harder questions"})
        st.dataframe(rows, use_container_width=True, hide_index=True)
        weakest = min(st.session_state.mastery.items(), key=lambda kv: kv[1]["correct"] / max(1, kv[1]["attempts"]))
        st.info(f"Suggested next step: revise **{weakest[0]}**, then retake a quiz. Mastery is currently a simple accuracy estimate, not a validated Bayesian learner model.")
    else:
        st.info("Complete quiz questions to start building your local session progress. Progress currently lasts only for this browser session.")

with sources_tab:
    st.subheader("Indexed source passages")
    st.caption("Citations are anchored to PDF page numbers, PowerPoint slide numbers, or transcript timestamps.")
    for src in sorted(set(c["source"] for c in chunks)):
        with st.expander(src):
            for c in [x for x in chunks if x["source"] == src]:
                st.markdown(f"**{c['location']}**")
                st.write(c["text"])
                st.divider()

st.divider()
st.caption("Honest MVP limitations: no direct video ingestion, no OCR for image-only pages/diagrams, no LLM synthesis, no persistent user database, and no external evaluation benchmark yet. These are next development tasks, not completed features.")
